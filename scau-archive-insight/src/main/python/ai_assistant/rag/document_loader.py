import os
import asyncio
import socket

from .url_security import (
    validate_browser_subresource_url,
    validate_public_http_url,
)


_NETWORK_API_BLOCK_SCRIPT = """
(() => {
    const blockedApis = [
        'WebSocket',
        'WebSocketStream',
        'EventSource',
        'Worker',
        'SharedWorker',
        'WebTransport',
        'RTCPeerConnection',
        'webkitRTCPeerConnection'
    ];
    for (const name of blockedApis) {
        const BlockedNetworkApi = class {
            constructor() {
                throw new Error(`${name} is disabled while importing knowledge URLs`);
            }
        };
        try {
            Object.defineProperty(globalThis, name, {
                configurable: false,
                writable: false,
                value: BlockedNetworkApi
            });
        } catch (_) {
            try { globalThis[name] = BlockedNetworkApi; } catch (_) {}
        }
    }
})();
"""


def load_pdf(path):
    """PyMuPDF 逐页提取文本"""
    import fitz
    doc = fitz.open(path)
    texts = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            texts.append(text)
    doc.close()
    return texts


def load_docx(path):
    """python-docx 提取段落"""
    from docx import Document
    doc = Document(path)
    return [p.text for p in doc.paragraphs if p.text.strip()]


def load_xlsx(path):
    """openpyxl 每行转文本"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    texts = []
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(str(c) for c in row if c is not None)
            if line.strip():
                texts.append(line)
    wb.close()
    return texts


def load_txt(path):
    """直接读取文本文件"""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    return [content]


async def _guard_browser_request(route, *, resolver, blocked_reasons):
    """阻止浏览器访问非公网 HTTP(S) 地址及其他网络协议。"""
    try:
        validate_browser_subresource_url(route.request.url, resolver=resolver)
    except ValueError as exc:
        blocked_reasons.append(str(exc))
        await route.abort()
        return
    await route.continue_()


async def load_url(url, *, resolver=socket.getaddrinfo):
    """用 Playwright + Chrome 渲染网页后提取文本，支持 JS 动态页面"""
    validated_url = validate_public_http_url(url, resolver=resolver)

    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = None
        blocked_reasons = []
        try:
            # 禁用 Service Worker，确保页面网络请求不会绕过路由拦截。
            context = await browser.new_context(service_workers="block")

            async def guard_request(route):
                await _guard_browser_request(
                    route,
                    resolver=resolver,
                    blocked_reasons=blocked_reasons,
                )

            await context.route("**/*", guard_request)
            # 禁用不受普通 HTTP 路由完整覆盖的长连接、Worker 和点对点网络 API。
            await context.add_init_script(_NETWORK_API_BLOCK_SCRIPT)

            page = await context.new_page()
            try:
                response = await page.goto(
                    validated_url,
                    wait_until="networkidle",
                    timeout=30000,
                )
            except Exception as exc:
                if blocked_reasons:
                    raise ValueError(
                        f"网页请求被安全策略阻止: {blocked_reasons[0]}"
                    ) from exc
                raise

            if blocked_reasons:
                raise ValueError(f"网页请求被安全策略阻止: {blocked_reasons[0]}")
            if response is not None and response.status >= 400:
                raise ValueError(f"网页请求失败，HTTP 状态码: {response.status}")

            # 再校验最终地址，覆盖浏览器重定向后的落点。
            validate_public_http_url(page.url, resolver=resolver)
            await page.wait_for_timeout(2000)
            if blocked_reasons:
                raise ValueError(f"网页请求被安全策略阻止: {blocked_reasons[0]}")
            text = await page.inner_text("body")
            return [text.strip()]
        finally:
            if context is not None:
                await context.close()
            await browser.close()


def load_document(file_path: str, file_type: str = None):
    """同步解析文档（PDF/DOCX/XLSX/TXT），返回文本列表"""
    if file_type is None:
        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        file_type = ext

    if file_type == "pdf":
        return load_pdf(file_path)
    elif file_type in ("docx", "doc"):
        return load_docx(file_path)
    elif file_type in ("xlsx", "xls"):
        return load_xlsx(file_path)
    elif file_type in ("txt", "csv"):
        return load_txt(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {file_type}")


async def load_document_async(file_path: str, file_type: str = None):
    """异步解析文档，目前只有 URL 需要异步"""
    if file_type is None:
        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        file_type = ext

    if file_type in ("html", "url"):
        return await load_url(file_path)
    else:
        return load_document(file_path, file_type)
